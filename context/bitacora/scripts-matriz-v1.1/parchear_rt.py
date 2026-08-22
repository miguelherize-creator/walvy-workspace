# -*- coding: utf-8 -*-
"""Aplica las respuestas de Kabeli al workbook de evidencia técnica.

Toca sólo las tres columnas de respuesta de `01_Requerimientos` (Estado,
Archivo/enlace, Comentarios) y agrega una hoja con el detalle. La hoja
`02_Detalle` no tiene columnas de respuesta —su «Comentarios / referencia» lo
escribió Walvy— así que no se modifica, igual que `00_Instrucciones` y
`03_Glosario`.
"""
import shutil
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from respuestas_rt import RESPUESTAS, RESUMENES

SRC, OUT = 'evidencia_base.xlsx', 'Walvy_Requerimientos_Evidencia_Tecnica_Kabeli_v1_0_respondido.xlsx'
shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT)

HOJA = next(n for n in wb.sheetnames if 'Requerimiento' in n)
ws = wb[HOJA]
HDR = 4
COL = {ws.cell(HDR, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(HDR, c).value}
c_estado, c_enlace, c_coment = COL['Estado'], COL['Archivo / enlace entregado'], COL['Comentarios']
INTOCABLES = [k for k in COL if COL[k] not in (c_estado, c_enlace, c_coment)]
filas = {ws.cell(r, 1).value: r for r in range(HDR + 1, ws.max_row + 1) if ws.cell(r, 1).value}
antes = {t: {k: ws.cell(filas[t], COL[k]).value for k in INTOCABLES} for t in filas}
altos_previos = {t: ws.row_dimensions[filas[t]].height for t in filas}

FUENTE = ws.cell(HDR + 1, 1).font.name or 'Calibri'
TAM = ws.cell(HDR + 1, 1).font.sz or 11

tocadas = 0
for rt, resp in RESPUESTAS.items():
    r = filas[rt]
    ws.cell(r, c_estado).value = resp['estado']
    ws.cell(r, c_enlace).value = resp['enlace']
    ws.cell(r, c_coment).value = RESUMENES[rt]
    tocadas += 3
    # Alto acotado: con el resumen en la celda ya no hace falta una fila enorme,
    # pero sí algo más que los 82 originales.
    ws.row_dimensions[r].height = 150

# --- hoja nueva con el detalle ---
NUEVA = '04_Respuesta Kabeli'
if NUEVA in wb.sheetnames:
    del wb[NUEVA]
ws4 = wb.create_sheet(NUEVA)
ws4['A1'] = 'Respuesta de Kabeli · detalle por requerimiento'
ws4['A1'].font = Font(name=FUENTE, sz=TAM + 3, bold=True)
ws4['A2'] = (
    'Detalle de lo que la hoja 01_Requerimientos resume. Un renglón por requerimiento. '
    'Se distingue lo que está en KabeliDev/main de lo que está en un PR abierto, y se declara '
    'que walvy-org/main está 99 commits atrás: hoy nada de lo referenciado es verificable en el '
    'repositorio que audita Walvy, y la sincronización es una decisión pendiente de Kabeli.'
)
ws4['A2'].font = Font(name=FUENTE, sz=TAM, italic=True)
ws4['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws4.merge_cells('A2:D2')
ws4.row_dimensions[2].height = 60

gris = PatternFill('solid', fgColor='D9D9D9')
borde = Border(*[Side(style='thin', color='BFBFBF')] * 4)
for i, t in enumerate(['ID', 'Estado', 'Archivo / enlace entregado', 'Detalle de la respuesta'], 1):
    c = ws4.cell(4, i, t)
    c.font = Font(name=FUENTE, sz=TAM, bold=True)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    c.fill, c.border = gris, borde

for i, rt in enumerate(sorted(RESPUESTAS), start=5):
    resp = RESPUESTAS[rt]
    for j, val in enumerate([rt, resp['estado'], resp['enlace'], resp['comentario']], 1):
        c = ws4.cell(i, j, val)
        c.font = Font(name=FUENTE, sz=TAM, bold=(j == 1))
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = borde
    # proporcional al largo del detalle, con techo para que la hoja siga navegable
    ws4.row_dimensions[i].height = min(520, 70 + len(resp['comentario']) // 11)

for col, w in zip('ABCD', [10, 46, 52, 120]):
    ws4.column_dimensions[col].width = w
ws4.freeze_panes = 'A5'
ws4.sheet_view.showGridLines = False

wb.save(OUT)
print(f'{tocadas} celdas escritas en {HOJA} · hoja "{NUEVA}" con {len(RESPUESTAS)} filas')

# ---------------- verificación ----------------
print('\n=== verificación ===')
v = openpyxl.load_workbook(OUT)
vs = v[HOJA]
print('1. hojas:', v.sheetnames)
alterados = [(t, k) for t in filas for k in INTOCABLES
             if vs.cell(filas[t], COL[k]).value != antes[t][k]]
print('2. columnas de Walvy intactas en las 9 filas:', not alterados, alterados or '')
print('3. filas sin responder:', [t for t in filas if t not in RESPUESTAS] or 'ninguna')
for hoja, orig in [('02_Detalle', 51), ('03_Glosario', 19), ('00_Instrucciones', 9)]:
    print(f'4. {hoja} intacta: {v[hoja].max_row == orig} ({v[hoja].max_row} filas)')
print('5. estados escritos:')
for rt in sorted(RESPUESTAS):
    print(f'     {rt}: {vs.cell(filas[rt], c_estado).value}')
vacias = [rt for rt in RESPUESTAS if not vs.cell(filas[rt], c_coment).value]
print('6. comentarios vacíos:', vacias or 'ninguno')
print('7. detalle en hoja 04:', all(v[NUEVA].cell(i, 4).value for i in range(5, 14)))
