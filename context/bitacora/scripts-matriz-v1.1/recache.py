# -*- coding: utf-8 -*-
"""Reinyecta los valores en caché de la columna R (Estado de cierre).

openpyxl no escribe valores cacheados, y sin LibreOffice no hay recálculo. La
salida se calcula replicando la fórmula del cliente y se VERIFICA contra la
caché del archivo original: I y N no se tocaron, así que el resultado tiene que
ser idéntico. Si difiere en una sola fila, se aborta.
"""
import re, shutil, zipfile
from xml.etree import ElementTree as ET
import openpyxl

BASE, OUT = 'base.xlsx', 'Walvy_M1_Matriz_Validacion_Tecnica_v1.1.xlsx'
HOJA, COL_R, FILAS = '01 Validación Técnica', 'R', range(5, 28)

def estado_de_cierre(I, N):
    """Réplica de la fórmula de la columna R del cliente."""
    I, N = (I or ''), (N or '')
    if I == 'No aplica': return 'No aplica'
    if I == 'Diferido aprobado': return 'Diferido – hito posterior'
    if I != 'Aceptado': return 'Pendiente definición alcance'
    if N == 'Conforme': return 'Conformado'
    if N == 'Observado': return 'Observado'
    if N == 'Evidencia pendiente': return 'Pendiente evidencia'
    if N == 'No aplica': return 'No aplica'
    if N == 'Diferido': return 'Diferido – hito posterior'
    return 'Pendiente cumplimiento'

wf = openpyxl.load_workbook(OUT)[HOJA]                      # I y N (fórmulas/literales)
cache_orig = openpyxl.load_workbook(BASE, data_only=True)[HOJA]
COLS = {wf.cell(4, c).value: c for c in range(1, wf.max_column + 1) if wf.cell(4, c).value}
cI, cN = COLS['ETAPA 1 · Validación alcance Walvy'], COLS['ETAPA 2 · Validación cumplimiento Walvy']

calculado, discrepa = {}, []
for r in FILAS:
    esperado = estado_de_cierre(wf.cell(r, cI).value, wf.cell(r, cN).value)
    original = cache_orig.cell(r, 18).value
    calculado[r] = esperado
    if esperado != original:
        discrepa.append((r, wf.cell(r, 1).value, original, esperado))

if discrepa:
    print('ABORTA — la réplica de la fórmula no reproduce la caché original:')
    for r, tid, o, e in discrepa:
        print(f'  fila {r} ({tid}): original={o!r} calculado={e!r}')
    raise SystemExit(1)
print(f'réplica de la fórmula validada contra la caché original en {len(FILAS)} filas')

# Resolver qué XML es la hoja 01
with zipfile.ZipFile(OUT) as z:
    wbx = ET.fromstring(z.read('xl/workbook.xml'))
    rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
NSm = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
NSr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
rid = {s.get('name'): s.get(f'{NSr}id') for s in wbx.iter(f'{NSm}sheet')}[HOJA]
target = {rl.get('Id'): rl.get('Target') for rl in rels}[rid]
target = target.lstrip('/')
path = target if target.startswith('xl/') else 'xl/' + target
print(f'hoja "{HOJA}" -> {path}')

def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

with zipfile.ZipFile(OUT) as z:
    xml = z.read(path).decode('utf-8')
    otros = [(i, z.read(i.filename)) for i in z.infolist() if i.filename != path]

inyectadas = 0
for r in FILAS:
    ref = f'{COL_R}{r}'
    # openpyxl deja  <c r="R7" s="2"><f>...</f><v /></c>  — el <v> va vacío.
    # Queda  <c r="R7" s="2" t="str"><f>...</f><v>Pendiente ...</v></c>
    pat = re.compile(
        r'<c r="' + ref + r'"([^>]*?)>(<f>.*?</f>)(<v\s*/>|<v></v>)?</c>', re.S)
    m = pat.search(xml)
    if not m:
        print(f'  AVISO: no se encontró celda con fórmula en {ref}'); continue
    attrs = m.group(1)
    attrs = attrs if ' t="str"' in attrs else attrs + ' t="str"'
    nuevo = f'<c r="{ref}"{attrs}>{m.group(2)}<v>{esc(calculado[r])}</v></c>'
    xml = xml[:m.start()] + nuevo + xml[m.end():]
    inyectadas += 1
print(f'valores en caché inyectados: {inyectadas}')

shutil.copy(OUT, OUT + '.tmp')
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as z:
    for info, data in otros:
        z.writestr(info, data)
    z.writestr(path, xml.encode('utf-8'))
import os; os.remove(OUT + '.tmp')
print('archivo reescrito')
